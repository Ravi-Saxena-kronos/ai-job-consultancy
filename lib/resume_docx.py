"""Build tailored resume as DOCX bytes."""

from __future__ import annotations

import re
from io import BytesIO


def _safe_filename_part(text: str, max_len: int = 40) -> str:
    s = re.sub(r"[^\w\-]+", "_", (text or "").strip())
    s = s.strip("_") or "Candidate"
    return s[:max_len]


def resume_docx_filename(seeker_name: str, company: str, app_id: str) -> str:
    name = _safe_filename_part(seeker_name.split()[0] if seeker_name else "Resume")
    co = _safe_filename_part(company, 30)
    aid = _safe_filename_part(app_id.replace("LNK-", "").replace("APP-", ""), 12)
    return f"Resume_{name}_{co}_{aid}.docx"


def build_tailored_docx(
    plain_text: str,
    *,
    seeker_name: str,
    job_title: str,
    company: str,
) -> bytes:
    from docx import Document

    doc = Document()
    heading = seeker_name.strip() or "Candidate"
    doc.add_heading(heading, level=0)
    if job_title:
        doc.add_paragraph(f"Applying for: {job_title}")
    doc.add_paragraph("")

    for block in plain_text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.endswith(":") and len(block) < 80:
            doc.add_heading(block.rstrip(":"), level=2)
        else:
            for line in block.split("\n"):
                line = line.strip()
                if line:
                    doc.add_paragraph(line)

    doc.add_paragraph("")
    doc.add_paragraph(
        f"Tailored for: {job_title} at {company}".strip(),
        style=None,
    )

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
